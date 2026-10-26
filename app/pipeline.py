"""Daily run: detect new Codal reports -> update the FIS workbook with the skill -> send it on Telegram.

Usage: python -m app.pipeline [--force] [--dry-run]
  --force     rebuild and send even if no new report was found
  --dry-run   build the workbook but do not call Claude, send to Telegram or save state
"""
import datetime
import fcntl
import json
import logging
import os
import re
import subprocess
import sys
import threading
import time
from pathlib import Path

from openpyxl import load_workbook

from . import claude_review, codal_watch, config, telegram
from .codal_watch import g2j, norm

log = logging.getLogger("pipeline")
SECRET_VARS = ("TELEGRAM_BOT_TOKEN", "ANTHROPIC_API_KEY", "TELEGRAM_CHAT_ID")
LATEST = config.OUTPUT_DIR / "FIS_ranking_latest.xlsx"
RESUMABLE = 75   # exit code: stopped early (Codal throttling, timeout); the scheduler retries


class Resumable(RuntimeError):
    """A step stopped early but kept its progress."""


# ---------------------------------------------------------------- state
def load_state():
    try:
        return json.loads(config.STATE_FILE.read_text(encoding="utf-8"))
    except FileNotFoundError:
        return {}


def save_state(state):
    tmp = config.STATE_FILE.with_suffix(".tmp")
    tmp.write_text(json.dumps(state, ensure_ascii=False, indent=1), encoding="utf-8")
    os.replace(tmp, config.STATE_FILE)


def write_json(path, obj):
    tmp = path.with_suffix(".tmp")
    tmp.write_text(json.dumps(obj, ensure_ascii=False, default=str), encoding="utf-8")
    os.replace(tmp, path)


# ---------------------------------------------------------------- skill steps
def run_skill(script, *args):
    """Run a skill batch script; secrets are not passed to it. Its output is logged live. Returns stdout."""
    env = {k: v for k, v in os.environ.items() if k not in SECRET_VARS}
    # -u: -I ignores PYTHONUNBUFFERED, and buffered output would only show when the step ends
    cmd = [sys.executable, "-I", "-u", str(Path(__file__).with_name("fis_run.py")),
           str(config.BATCH_DIR / script), *map(str, args)]
    log.info("running %s", script)
    t0 = time.monotonic()
    with subprocess.Popen(cmd, env=env, stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True) as p:
        timer = threading.Timer(config.STEP_TIMEOUT, p.kill)
        timer.start()
        lines = []
        for line in p.stdout:
            line = line.rstrip("\n")
            lines.append(line)
            log.info("  [%s] %s", script, line)
        p.wait()
        timer.cancel()
    elapsed = time.monotonic() - t0
    if p.returncode == RESUMABLE:
        raise Resumable(f"{script}: {lines[-1] if lines else 'stopped'}")
    if p.returncode != 0:
        if elapsed >= config.STEP_TIMEOUT:  # the batch scripts save each symbol, so this resumes too
            raise Resumable(f"{script} timed out after {_dur(elapsed)}")
        raise RuntimeError(f"{script} exited {p.returncode}: " + "\n".join(lines[-30:])[-1500:])
    log.info("%s finished in %s", script, _dur(elapsed))
    return "\n".join(lines)


def _dur(seconds):
    m, s = divmod(int(seconds), 60)
    h, m = divmod(m, 60)
    return f"{h}h{m:02d}m{s:02d}s" if h else f"{m}m{s:02d}s"


def invalidate(new):
    """Mark collected data of symbols with new letters for re-collection.

    results/ files are deleted (run_all.py collects missing ones). investees/ files are kept
    but marked stale: holdings.py only re-collects investees that are still in a holding's
    portfolio, and a deleted file of a company that left every portfolio would never come back.
    Returns (affected {key: name}, stale investee files).
    """
    affected, stale = {}, set()
    for n in new:
        t = n["targets"]
        if t["results"]:
            (config.WORK_DIR / "results" / f"{t['results']}.json").unlink(missing_ok=True)
        if t["investee"]:
            stale.add(config.WORK_DIR / "investees" / f"{t['investee']}.json")
        affected[n["key"]] = t["results"] or t.get("add") or t["investee"]
    return affected, stale


def load_symbols():
    return json.load(open(config.WORK_DIR / "symbols.json", encoding="utf-8"))


def add_symbols(names):
    """Append symbols to symbols.json (skipping ones already there). Returns the added ones."""
    symbols = load_symbols()
    have = {norm(s) for s in symbols}
    added = []
    for s in names:
        if norm(s) not in have:
            have.add(norm(s))
            added.append(s)
    if added:
        backup = config.WORK_DIR / "symbols.user.json"   # the original hand-made list, kept once
        if not backup.exists():
            backup.write_text(json.dumps(symbols, ensure_ascii=False), encoding="utf-8")
        write_json(config.WORK_DIR / "symbols.json", symbols + added)
    return added


def backfill_candidates(n):
    """Up to n listed companies that are not in the workbook yet (so it covers the whole market)."""
    if n <= 0:
        return []
    known = {norm(s) for s in load_symbols()}
    known |= {norm(f.stem) for f in (config.WORK_DIR / "investees").glob("*.json")}
    listed = codal_watch.listed_stocks()
    return [listed[k] for k in sorted(listed) if k not in known][:n]


def pending_symbols():
    """Symbols run_all.py will collect (same rule as its needs_collect)."""
    out = []
    for s in load_symbols():
        f = config.WORK_DIR / "results" / f"{s}.json"
        try:
            if "fetch failed" not in (json.loads(f.read_text(encoding="utf-8")).get("error") or ""):
                continue
        except (FileNotFoundError, ValueError):
            pass
        out.append(s)
    return out


def refresh_prices(stale=()):
    """Today's closing prices for every collected symbol (one market-watch request).

    Also marks unchanged investees as collected today, so holdings.py reuses them instead of
    re-collecting a few hundred companies from Codal. Stale investees get an old mtime so
    holdings.py re-collects them.
    """
    from fetch import market_watch
    log.info("refreshing prices from TSETMC")
    mw = market_watch()
    n = 0
    for sub in ("results", "investees"):
        d = config.WORK_DIR / sub
        if not d.is_dir():
            continue
        for f in d.glob("*.json"):
            r = json.loads(f.read_text(encoding="utf-8"))
            m = mw.get(norm(r.get("symbol")))
            if m and m.get("price") and not r.get("error"):
                r["price"] = m["price"]
                n += 1
            write_json(f, r)
            if f in stale:
                os.utime(f, (0, 0))
    log.info("prices refreshed for %d symbols (%d investees marked for re-collection)", n, len(stale))


def build_excel():
    today = g2j(datetime.date.today()).replace("/", "-")
    out = config.OUTPUT_DIR / f"FIS_ranking_{today}.xlsx"
    log.info("building workbook %s", out.name)
    stdout = run_skill("build_excel.py", config.WORK_DIR, out)
    m = re.search(r"ranked=(\d+) excluded=(\d+) avg=([\d.]+)", stdout)
    stats = {"ranked": int(m[1]), "excluded": int(m[2]), "avg": float(m[3])} if m else {}
    tmp = LATEST.with_suffix(".tmp")
    tmp.write_bytes(out.read_bytes())
    os.replace(tmp, LATEST)
    dated = sorted(config.OUTPUT_DIR.glob("FIS_ranking_1*.xlsx"))
    for old in dated[:-config.KEEP_EXCEL_FILES]:
        log.info("removing old workbook %s", old.name)
        old.unlink()
    log.info("workbook ready: %s", stats or "no stats line in build_excel output")
    return out, stats


# ---------------------------------------------------------------- comparison
def read_ranking(path):
    """{norm symbol: row} from the FIS and holdings sheets of a workbook built by the skill."""
    if not path.exists():
        return {}, {}
    wb = load_workbook(path, read_only=True)
    fis, hold = {}, {}
    rows = list(wb.worksheets[0].iter_rows(values_only=True))
    col = {h: i for i, h in reversed(list(enumerate(rows[0] if rows else [])))}  # first column per header
    get = lambda r, h: r[col[h]] if h in col else None  # noqa: E731
    for r in rows[1:]:
        if get(r, "نماد"):
            fis[norm(get(r, "نماد"))] = {
                "symbol": get(r, "نماد"), "rank": get(r, "رتبه"), "score": get(r, "نمره FIS (نرمال‌شده از ۱۰۰)"),
                "raw": get(r, "امتیاز خام"), "above_avg": get(r, "بالاتر از میانگین گروه؟"),
                "pe_status": get(r, "وضعیت P/E (اخزا ۴۰٪)"), "fwd_pe": get(r, "P/E آینده‌نگر"),
                "fwd_pd": get(r, "P/D آینده‌نگر")}
    if len(wb.worksheets) > 1:
        for r in list(wb.worksheets[1].iter_rows(min_row=2, values_only=True)):
            if r[1]:
                hold[norm(r[1])] = {"symbol": r[1], "rank": r[0], "p_nav80": r[3]}
    wb.close()
    return fis, hold


def result_issues(name):
    for sub in ("results", "investees"):
        f = config.WORK_DIR / sub / f"{name}.json"
        if f.exists():
            r = json.loads(f.read_text(encoding="utf-8"))
            return {"error": r.get("error"), "notes": r.get("notes"), "report": r.get("report"),
                    "monthly": r.get("m_title")}
    return {"error": "داده جمع‌آوری نشد"}


def build_facts(new, affected, before, after, stats):
    (fb, hb), (fa, ha) = before, after
    changes = []
    for key, name in affected.items():
        changes.append({"symbol": name, "fis_before": fb.get(key), "fis_after": fa.get(key),
                        "holding_before": hb.get(key), "holding_after": ha.get(key),
                        "data": result_issues(name)})
    top = sorted(fa.values(), key=lambda x: x["rank"] or 1e9)[:15]
    return {"date": g2j(datetime.date.today()), "new_letters":
            [{k: n[k] for k in ("symbol", "type", "title", "published")} for n in new],
            "changes": changes, "top15": top, "stats": stats}


def plain_summary(facts):
    lines = [f"به‌روزرسانی FIS — {facts['date']}", ""]
    if facts["new_letters"]:
        lines.append("گزارش‌های جدید:")
        lines += [f"• {n['symbol']}: {n['title']}" for n in facts["new_letters"]]
        lines.append("")
    lines.append("تغییر نمره:")
    for c in facts["changes"]:
        b, a = c["fis_before"] or {}, c["fis_after"] or {}
        if c["data"].get("error"):
            lines.append(f"• {c['symbol']}: ⚠️ {c['data']['error']}")
        elif a or b:
            lines.append(f"• {c['symbol']}: {b.get('score', '—')} (رتبه {b.get('rank', '—')}) ← "
                         f"{a.get('score', '—')} (رتبه {a.get('rank', '—')})")
        elif c["holding_after"]:
            lines.append(f"• {c['symbol']}: P/NAV80 = {c['holding_after']['p_nav80']}٪")
        else:
            lines.append(f"• {c['symbol']}: در جدول FIS نیست (هلدینگ زیرمجموعه)")
    if facts.get("added_from_reports"):
        lines += ["", "شرکت‌های تازه اضافه‌شده به اکسل (گزارش جدید): " + "، ".join(facts["added_from_reports"])]
    if facts.get("backfilled"):
        lines += ["", f"تکمیل تدریجی بازار: {len(facts['backfilled'])} نماد جدید — " + "، ".join(facts["backfilled"])]
    s = facts["stats"]
    if s:
        lines += ["", f"رتبه‌بندی‌شده: {s['ranked']} | حذف‌شده: {s['excluded']} | میانگین: {s['avg']}"]
    lines += ["", "۱۰ نماد اول: " + "، ".join(f"{x['symbol']} ({x['score']})" for x in facts["top15"][:10]),
              "", "این خروجی پژوهشی است و توصیه سرمایه‌گذاری نیست."]
    return "\n".join(lines)


# ---------------------------------------------------------------- main
def run(force=False, dry_run=False):
    """One run. Progress is saved in state.json as it goes: if Codal throttles us, the next run
    continues where this one stopped (collected symbols are kept, new letters are not lost)."""
    t0 = time.monotonic()
    log.info("run started (force=%s, dry_run=%s)", force, dry_run)
    state = load_state()
    today = datetime.date.today().isoformat()
    if state.get("pending_new") or state.get("pending_affected"):
        log.info("resuming an unfinished run: %d letters not reported yet", len(state.get("pending_new", [])))
    log.info("[1/6] searching Codal for new letters (since %s)", codal_watch.from_date(state))
    fresh = codal_watch.find_new(state)
    for n in fresh:
        log.info("  new %s: %s — %s", n["type"], n["symbol"], n["title"])
    backfill = backfill_candidates(config.BACKFILL_PER_DAY) if state.get("backfill_date") != today else []
    new = state.get("pending_new", []) + fresh
    pending = pending_symbols()
    log.info("new letters: %d (%d this run), symbols without data: %d, backfill: %d",
             len(new), len(fresh), len(pending), len(backfill))
    if not new and not pending and not backfill and not force:
        if dry_run:
            return
        log.info("nothing to update")
        if config.SEND_WHEN_NO_CHANGE:
            log.info("sending 'no new report' message to Telegram")
            telegram.send_text(f"FIS — {g2j(datetime.date.today())}: گزارش جدیدی در کدال منتشر نشده است.")
        state["last_success"] = today
        save_state(state)
        return

    before = read_ranking(LATEST)
    affected, stale = invalidate(fresh)
    affected = {**state.get("pending_affected", {}), **affected}
    added = add_symbols([n["targets"]["add"] for n in fresh if n["targets"].get("add")])
    for s in pending_symbols():      # user-added symbols and companies with a first report
        affected.setdefault(norm(s), s)
    backfill = add_symbols(backfill)
    log.info("added to symbols.json: %d from new reports, %d backfill", len(added), len(backfill))
    # saved before the long steps, so a stopped run neither loses these letters nor
    # invalidates (and re-collects) the same symbols again next time
    state.update(
        pending_new=new, pending_affected=affected,
        pending_added=state.get("pending_added", []) + added,
        pending_backfill=state.get("pending_backfill", []) + backfill,
        stale_investees=sorted(set(state.get("stale_investees", [])) | {f.stem for f in stale}),
        seen=(state.get("seen", []) + [t for n in fresh for t in n["tracing_nos"]])[-5000:],
    )
    if config.BACKFILL_PER_DAY > 0 and state.get("backfill_date") != today:
        state["backfill_date"] = today
    if not dry_run:
        save_state(state)

    log.info("[2/6] collecting %d symbols from Codal (~1 min each)", len(pending_symbols()))
    run_skill("run_all.py", config.WORK_DIR)
    run_skill("recompute.py", config.WORK_DIR)  # results saved by an older method, re-derived from the cache
    log.info("[3/6] refreshing prices")
    inv = config.WORK_DIR / "investees"
    refresh_prices({inv / f"{name}.json" for name in state["stale_investees"]})
    if config.RUN_HOLDINGS:
        log.info("[4/6] evaluating holdings")
        try:
            run_skill("holdings.py", config.WORK_DIR)
        finally:  # investees re-collected before a stop are fresh now
            state["stale_investees"] = [n for n in state["stale_investees"]
                                        if (inv / f"{n}.json").exists() and (inv / f"{n}.json").stat().st_mtime == 0]
            if not dry_run:
                save_state(state)
    else:
        log.info("[4/6] holdings skipped (RUN_HOLDINGS=false)")
    log.info("[5/6] building the ranking workbook")
    out, stats = build_excel()
    facts = build_facts(new, affected, before, read_ranking(out), stats)
    facts["added_from_reports"], facts["backfilled"] = state["pending_added"], state["pending_backfill"]

    if dry_run:
        log.info("dry run, summary would be:\n%s", plain_summary(facts))
        return
    log.info("[6/6] writing the summary and sending to Telegram")
    summary = claude_review.summarize(facts)
    if not summary:
        log.info("using the plain summary (Claude not configured or failed)")
        summary = plain_summary(facts)
    syms = "، ".join(dict.fromkeys(n["symbol"] for n in new))[:600] or "بدون گزارش جدید"
    log.info("telegram: sending %s", out.name)
    telegram.send_document(out, caption=f"📊 رتبه‌بندی FIS — {facts['date']}\nگزارش جدید: {syms}")
    log.info("telegram: sending summary (%d chars)", len(summary))
    telegram.send_text(summary)

    state["last_success"] = today
    for k in ("pending_new", "pending_affected", "pending_added", "pending_backfill"):
        state.pop(k, None)
    save_state(state)
    log.info("done: %s in %s", out.name, _dur(time.monotonic() - t0))


def main(argv=None):
    argv = sys.argv[1:] if argv is None else argv
    logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(name)s: %(message)s")
    logging.getLogger("httpx").setLevel(logging.WARNING)   # request lines would include URLs
    config.OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    with open(config.LOCK_FILE, "w") as lock:
        try:
            fcntl.flock(lock, fcntl.LOCK_EX | fcntl.LOCK_NB)
        except BlockingIOError:
            log.warning("another run is in progress; skipping")
            return 0
        try:
            run(force="--force" in argv, dry_run="--dry-run" in argv)
            return 0
        except (Resumable, codal_watch.Throttled) as e:
            log.warning("run stopped, progress is saved and the next run continues: %s", e)
            return RESUMABLE
        except Exception as e:  # noqa: BLE001
            log.exception("run failed")
            if "--dry-run" in argv:
                return 1
            try:
                telegram.send_text(f"⚠️ اجرای روزانه FIS ناموفق بود:\n{type(e).__name__}: {str(e)[:1500]}")
            except Exception:  # noqa: BLE001
                log.exception("could not report the failure on Telegram")
            return 1


if __name__ == "__main__":
    sys.exit(main())
