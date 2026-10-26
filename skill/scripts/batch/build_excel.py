import glob
import json
import os
import sys

"""Score collected results with fis_score.py and write the ranking workbook.

Usage: python3 -I build_excel.py WORKDIR OUT.xlsx [PYLIB_DIR]
PYLIB_DIR: folder where openpyxl was installed with `pip install --target` (if not
installed system-wide).
"""
HERE = os.path.dirname(os.path.abspath(__file__))
WORK = os.path.abspath(sys.argv[1]) if len(sys.argv) > 1 else os.getcwd()
if len(sys.argv) > 3:
    sys.path.insert(0, os.path.abspath(sys.argv[3]))
sys.path.insert(0, os.path.join(HERE, ".."))
sys.path.insert(0, HERE)
from collect import forward_net  # noqa: E402
from fis_score import ORDER, THRESHOLDS, evaluate  # noqa: E402
from openpyxl import Workbook  # noqa: E402
from openpyxl.formatting.rule import ColorScaleRule  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

RF = 40
OUT = sys.argv[2] if len(sys.argv) > 2 else os.path.join(WORK, "FIS_ranking.xlsx")
symbols = json.load(open(os.path.join(WORK, "symbols.json"), encoding="utf-8"))
raw = {}
for s in symbols:
    p = os.path.join(WORK, "results", f"{s}.json")
    if os.path.exists(p):
        raw[s] = json.load(open(p, encoding="utf-8"))


def pos(x):
    return x is not None and x > 0


def pct(a, b):
    """a / b in percent, or None (b must be positive: ROE of negative equity means nothing)."""
    return round(a / b * 100, 1) if a is not None and pos(b) else None


def to_input(r):
    mcap = r["price"] * r["shares"] / 1e6 if r.get("price") and r.get("shares") else None
    # cash raised in a capital increase after the report has not earned anything in the reports yet:
    # valuation ratios use the market cap without it (= the market cap before the increase)
    val_mcap = mcap - (r.get("cash_increase") or 0) if mcap else None
    ttm = r.get("ttm_net")
    loss = (ttm <= 0) if ttm is not None else (r.get("net") is not None and r["net"] <= 0)
    rev, cogs = r.get("rev"), r.get("cogs")
    d = {
        "company": r["symbol"],
        "is_loss_making": loss,
        "has_monthly_reports": True,
        "market_cap": val_mcap,
        "months_elapsed": r.get("months"),
        "ytd_net_income_fis": r.get("fis"),
        "ytd_sales": rev if pos(rev) else None,
        "ttm_net_income": ttm,
        "sales_period_prior_year": r.get("rev_prev_same") if pos(r.get("rev_prev_same")) else None,
        # turnover periods: last-12-month revenue / cost of sales over average balances, 365 days
        "period_days": 365 if pos(r.get("ttm_rev")) else r.get("period_days", 365),
        "period_sales": r["ttm_rev"] if pos(r.get("ttm_rev")) else (rev if pos(rev) else None),
        "period_cogs": (r["ttm_cogs"] if pos(r.get("ttm_rev")) and pos(r.get("ttm_cogs"))
                        else (cogs if pos(cogs) and not pos(r.get("ttm_rev")) else None)),
        "risk_free_rate": RF,
        "balance_sheet": {
            "receivables": r.get("receivables") if pos(rev) else None,
            "receivables_avg": r.get("avg_receivables") if pos(rev) else None,
            "total_assets": r.get("total_assets"),
            "average_inventory": r.get("avg_inventory") if pos(cogs) else None,
        },
    }
    # Paper: forward sales/earnings from monthly reports (avg monthly x 12). Monthly reports give
    # sales only, so forward earnings = forward sales x net margin of the latest interim statement,
    # without non-operating gains (collect.forward_net).
    if r.get("m_cum_rev") and r.get("m_months"):
        fwd_net, fwd_sales = forward_net(r)
        d["forward_sales"] = fwd_sales
        d["months_elapsed"] = r["m_months"]
        if fwd_net is not None:
            d["forward_net_income"] = fwd_net
        elif r.get("fis") is not None and r.get("months"):
            d["forward_net_income"] = r["fis"] * 12 / r["months"]
        if r.get("m_prev_rev") and r["m_prev_rev"] > 0:
            d["sales_period_current"] = r["m_cum_rev"]
            d["sales_period_prior_year"] = r["m_prev_rev"]
        else:
            d["sales_period_current"] = rev if pos(rev) else None
    else:
        d["sales_period_current"] = rev if pos(rev) else None
    qc, qp = r.get("q_cur") or [None, None], r.get("q_prev") or [None, None]
    q = {}
    if pos(qc[0]) and qc[1] is not None:
        q["current"] = {"revenue": qc[0], "fis_profit": qc[1]}
    if pos(qp[0]) and qp[1] is not None:
        q["previous"] = {"revenue": qp[0], "fis_profit": qp[1]}
    d["quarters"] = q
    dv = r.get("div_ratios") or []
    if dv:
        d["dividends"] = {"last_year": {"dividend": dv[0][0], "retained_earnings": dv[0][1]}}
        if len(dv) > 1:
            d["dividends"]["this_year"] = {"dividend": dv[1][0], "retained_earnings": dv[1][1]}
    return d, val_mcap


def _n(sym):  # TSETMC uses Arabic ي/ك, the user's list Persian ی/ک
    return str(sym).replace("ي", "ی").replace("ك", "ک").replace(" ", "").replace("\u200c", "")


# subsidiaries of holdings (collected by holdings.py) are scored on the FIS tab too
parents = {}
hpath0 = os.path.join(WORK, "holdings.json")
if os.path.exists(hpath0):
    for h in json.load(open(hpath0, encoding="utf-8")):
        for d in h.get("detail", []):
            if d.get("symbol"):
                parents.setdefault(_n(d["symbol"]), []).append(f"{h['symbol']} ({d['own'] * 100:.1f}٪)")
inv_dir = os.path.join(WORK, "investees")
if os.path.isdir(inv_dir):
    for f in sorted(os.listdir(inv_dir)):
        r = json.load(open(os.path.join(inv_dir, f), encoding="utf-8"))
        if r.get("symbol") and _n(r["symbol"]) not in {_n(k) for k in raw} and not r.get("error"):
            r["_from_holding"] = True
            raw[r["symbol"]] = r


def is_holding(r):
    f = (r.get("format") or "").replace("\u200c", " ")
    return not r.get("error") and "سرمایه گذاری" in f and "صندوق" not in f


MAX_RECEIVABLES_TO_ASSETS = 80  # % ; holdings are exempt (their receivables are mostly dividends due)
rows, excluded = [], []
for s, r in raw.items():
    if r.get("error"):
        excluded.append((s, r.get("name"), r["error"], r.get("report_url")))
        continue
    if is_holding(r):  # scored separately on tab 2 (NAV / look-through method)
        continue
    inp, mcap = to_input(r)
    o = evaluate(inp)
    if not o["eligible"]:
        excluded.append((s, r.get("name"), "زیان‌ده (سود ۱۲ ماه گذشته ≤ صفر) — طبق مقاله حذف", r.get("report_url")))
        continue
    vals = {x["key"]: x for x in o["rows"]}
    rta = vals["receivables_to_assets"]["value"]
    if isinstance(rta, (int, float)) and rta > MAX_RECEIVABLES_TO_ASSETS:
        excluded.append((s, r.get("name"), f"مطالبات {rta:.0f}٪ کل دارایی‌هاست (بیش از {MAX_RECEIVABLES_TO_ASSETS}٪) — "
                         "بخش بزرگی از دارایی وصول‌نشده است و کیفیت سود قابل اتکا نیست", r.get("report_url")))
        continue
    # P/D = forward P/E / payout ratio (forward dividend = forward earnings x the payout variable)
    pe, payout = vals["forward_pe"]["value"], vals["dividend_payout"]["value"]
    pd_fwd = pe / (payout / 100) if isinstance(pe, (int, float)) and pe > 0 and payout else None
    rows.append({"r": r, "o": o, "mcap": mcap, "vals": vals, "pd_fwd": pd_fwd,
                 "pe_ttm": (mcap / r["ttm_net"]) if mcap and r.get("ttm_net") else None})

# a normalized score from few computable variables (banks, insurers, funds-like companies) is not
# comparable with the rest: those go to the «سایر شرکت‌ها» sheet, outside the ranking and the average
MIN_MAX_AVAILABLE = 70
others = [x for x in rows if x["o"]["max_available"] < MIN_MAX_AVAILABLE]
rows = [x for x in rows if x["o"]["max_available"] >= MIN_MAX_AVAILABLE]
avg = sum(x["o"]["normalized_100"] for x in rows) / len(rows) if rows else 0
rows.sort(key=lambda x: (-x["o"]["normalized_100"], -x["o"]["total"]))
others.sort(key=lambda x: (-x["o"]["normalized_100"], -x["o"]["total"]))

wb = Workbook()
HFILL = PatternFill("solid", fgColor="1F4E78")
HFONT = Font(bold=True, color="FFFFFF")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)


def header(ws, cols, widths=None):
    ws.append(cols)
    for i, c in enumerate(ws[1], 1):
        c.fill, c.font, c.alignment = HFILL, HFONT, CENTER
        ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(i, 13)
    ws.row_dimensions[1].height = 45
    ws.sheet_view.rightToLeft = True
    ws.freeze_panes = "D2"


ws = wb.active
ws.title = "رتبه‌بندی FIS"
# (key, header): main ratios first, then the other fields, then the FIS score of each variable.
# Keys «v:<var>» are a variable's value, «s:<var>» its score.
LAYOUT = [("rank", "رتبه"), ("symbol", "نماد"), ("name", "نام شرکت"),
          ("score", "نمره FIS (نرمال‌شده از ۱۰۰)"), ("raw", "امتیاز خام"), ("max", "حداکثر قابل‌محاسبه"),
          ("pe_signal", "وضعیت P/E (اخزا ۴۰٪)"), ("v:forward_pe", None), ("pd_fwd", "P/D آینده‌نگر"),
          ("v:forward_ps", None), ("v:receivables_to_assets", None), ("v:collection_days", None),
          ("v:inventory_days", None), ("v:dividend_payout", None), ("v:fis_margin", None),
          ("v:margin_trend", None), ("v:sales_growth", None),
          ("recurring", "سود تکرارشونده افزوده به برآورد (سالانه، میلیارد تومان)"),
          ("gains", "سایر درآمدها و سود غیرعملیاتی دوره (میلیارد تومان)"), ("gains_pct", "سهم آن از سود خالص (%)"),
          ("above", "بالاتر از میانگین گروه؟"), ("price", "قیمت (ریال)"), ("mcap", "ارزش بازار (میلیارد تومان)"),
          ("cash_inc", "آورده نقدی افزایش سرمایه بعد از گزارش (میلیارد تومان)"), ("report", "آخرین گزارش"),
          ("months", "ماه‌های دوره"), ("pe_ttm", "P/E TTM"), ("roa", "ROA (%)"), ("roe", "ROE (%)"),
          ("v:earnings_expectation_gap", None), ("v:fiscal_year_coverage", None),
          ("parents", "زیرمجموعه هلدینگ (درصد مالکیت)"), ("source", "منبع نماد"),
          ("missing", "متغیرهای فاقد داده"), ("notes", "یادداشت"), ("link", "لینک کدال")]
LAYOUT += [(f"s:{k}", "امتیاز {} از {}".format(
    THRESHOLDS[k]["label"].replace(" (%)", "").replace(" (روز)", ""),
    str(THRESHOLDS[k]["max"]).translate(str.maketrans("0123456789", "۰۱۲۳۴۵۶۷۸۹")))) for k in ORDER]
LAYOUT = [(k, h or THRESHOLDS[k[2:]]["label"]) for k, h in LAYOUT]
assert {k[2:] for k, _ in LAYOUT if k.startswith("v:")} == set(ORDER)
KEYS = [k for k, _ in LAYOUT]
cols = [h for _, h in LAYOUT]
LINK = KEYS.index("link")
WIDTHS = {KEYS.index(k) + 1: w for k, w in (("symbol", 11), ("name", 26), ("pe_signal", 22), ("cash_inc", 34),
                                              ("missing", 40), ("notes", 40), ("link", 30))}


def fis_line(x, rank):
    r, o = x["r"], x["o"]
    line = {
        "rank": rank, "symbol": r["symbol"], "name": r.get("name"), "score": o["normalized_100"],
        "raw": o["total"], "max": o["max_available"], "pe_signal": o.get("pe_signal", "—"),
        "above": ("بله" if o["normalized_100"] > avg else "خیر") if rank != "—" else "—", "price": r.get("price"),
        "mcap": round((x["mcap"] + (r.get("cash_increase") or 0)) / 1e4, 1) if x["mcap"] else None,
        "cash_inc": round(r["cash_increase"] / 1e4, 1) if r.get("cash_increase") else None,
        "report": r.get("report"), "months": r.get("months"),
        "pe_ttm": round(x["pe_ttm"], 2) if x["pe_ttm"] else None,
        "pd_fwd": round(x["pd_fwd"], 2) if x["pd_fwd"] else None,
        "roa": pct(r.get("ttm_net"), r.get("total_assets")), "roe": pct(r.get("ttm_net"), r.get("equity")),
        "gains": round(r["gains"] / 1e4, 1) if r.get("gains") is not None else None,
        "gains_pct": pct(r.get("gains"), r.get("net")),
        "recurring": round(r["recurring_gains"] / 1e4, 1) if r.get("recurring_gains") else None,
        "parents": "، ".join(parents.get(_n(r["symbol"]), [])),
        "source": "زیرمجموعه هلدینگ" if r.get("_from_holding") else "فهرست شما",
        "missing": "، ".join(o["missing"]), "notes": " / ".join(r.get("notes", []) + o.get("notes", [])),
        "link": r.get("report_url"),
    }
    for k in ORDER:
        v = x["vals"][k]
        val = round(v["value"], 2) if isinstance(v["value"], (int, float)) else v["value"]
        if k == "forward_pe" and val is not None and val < 0:
            val = "منفی (سود آینده‌نگر ≤ ۰)"
        line[f"v:{k}"], line[f"s:{k}"] = val, v["score"]
    return [line[k] for k in KEYS]


def fis_sheet(ws, items, ranked):
    header(ws, cols, WIDTHS)
    for i, x in enumerate(items, 1):
        ws.append(fis_line(x, i if ranked else "—"))
    n = len(items) + 1
    ws.conditional_formatting.add(f"D2:D{n}", ColorScaleRule(start_type="min", start_color="F8696B",
                                                              mid_type="percentile", mid_value=50, mid_color="FFEB84",
                                                              end_type="max", end_color="63BE7B"))
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(horizontal="center", vertical="center")
        if row[LINK].value:
            row[LINK].hyperlink = row[LINK].value
            row[LINK].style = "Hyperlink"


fis_sheet(ws, rows, ranked=True)
fis_sheet(wb.create_sheet("سایر شرکت‌ها"), others, ranked=False)  # ends up after the holdings sheets

# ---------------------------------------------------------------- tab 2: holdings
hpath = os.path.join(WORK, "holdings.json")
hold = json.load(open(hpath, encoding="utf-8")) if os.path.exists(hpath) else []
ok = sorted([h for h in hold if not h.get("error") and h.get("p_nav80") is not None], key=lambda h: h["p_nav80"])
def _hnotes(h):
    n = list(h.get("notes", []))
    mc = h.get("market_cap") or 0
    if h.get("p_nav80") and h["p_nav80"] > 1.5 and h.get("listed_value_today", 0) < 0.5 * mc:
        n.insert(0, "بیشتر ارزش بازار پشتوانه سهام بورسی ندارد (دارایی غیربورسی/سایر که طبق روش حذف شده)؛ "
                    "نسبت بالا لزوماً به معنی گرانی نیست")
    if h.get("pe_fwd") and h["pe_fwd"] > 50:
        n.append("سود آینده‌نگر از سهام بورسی کوچک است (مجامع سرمایه‌پذیرها خارج از سال مالی یا پورتفوی بورسی کوچک)")
    return n


wh = wb.create_sheet("سرمایه‌گذاری و هلدینگ‌ها", 1)
B = 1e4  # million rials -> billion tomans
hcols = ["رتبه", "نماد", "نام", "ارزش بازار ÷ (۸۰٪ NAV) (%)", "قیمت (ریال)", "ارزش بازار (میلیارد تومان)",
         "NAV (میلیارد تومان)", "حقوق صاحبان سهام (میلیارد تومان)", "سود تحقق‌نیافته پورتفوی بورسی (میلیارد تومان)",
         "ارزش روز سهام بورسی (میلیارد تومان)", "بهای تمام شده سهام بورسی (میلیارد تومان)",
         "ارزش همان سهام در گزارش (میلیارد تومان)", "بدهی‌ها (میلیارد تومان)", "ROA (%)", "ROE (%)",
         "P/D آینده‌نگر", "سود تقسیمی آینده‌نگر هر سهم (ریال)", "P/D آخرین مجمع", "سود تقسیمی آخرین مجمع هر سهم (ریال)",
         "P/E آینده‌نگر", "P/E TTM", "P/B", "P/A", "نسبت بدهی",
         "سود آینده‌نگر (میلیارد تومان)", "سود نقدی دریافتی آینده‌نگر (میلیارد تومان)", "هزینه‌های هلدینگ (میلیارد تومان)",
         "نسبت تقسیم سود هلدینگ", "سال مالی مبنا", "تعداد سهام بورسی پورتفو", "گزارش پورتفو", "یادداشت", "لینک کدال"]
header(wh, hcols, {2: 11, 3: 26, 28: 24, 30: 34, 31: 45, 32: 30})
r2 = lambda v, n=2: round(v, n) if isinstance(v, (int, float)) else v  # noqa: E731
for i, h in enumerate(ok, 1):
    wh.append([i, h["symbol"], h.get("name"), r2(h["p_nav80"] * 100, 1), h.get("price"), r2((h.get("market_cap") or 0) / B, 1),
               r2(h["nav"] / B, 1), r2((h.get("equity") or 0) / B, 1), r2((h.get("unrealized") or 0) / B, 1),
               r2(h["listed_value_today"] / B, 1), r2((h.get("listed_cost") or 0) / B, 1),
               r2(h["listed_value_report"] / B, 1), r2((h.get("total_liabilities") or 0) / B, 1),
               r2(h["roa"] * 100, 1) if h.get("roa") is not None else None,
               r2(h["roe"] * 100, 1) if h.get("roe") is not None else None,
               r2(h.get("pd_fwd")), r2(h.get("dps_fwd"), 0), r2(h.get("pd_last")), r2(h.get("dps_last"), 0),
               r2(h.get("pe_fwd")), r2(h.get("pe_ttm")), r2(h.get("pb")), r2(h.get("pa")), r2(h.get("debt_ratio"), 3),
               r2(h["net_fwd"] / B, 1), r2(h["cash_net_fwd"] / B, 1), r2(h["costs"] / B, 1), r2(h.get("payout"), 3),
               h.get("fy_window"), h.get("n_listed"), h.get("portfolio_report"), " / ".join(_hnotes(h)),
               h.get("report_url")])
nh = len(ok) + 1
if ok:
    wh.conditional_formatting.add(f"D2:D{nh}", ColorScaleRule(start_type="min", start_color="63BE7B",
                                                               mid_type="percentile", mid_value=50, mid_color="FFEB84",
                                                               end_type="max", end_color="F8696B"))
for row in wh.iter_rows(min_row=2):
    for c in row:
        c.alignment = Alignment(horizontal="center", vertical="center")
    if row[-1].value:
        row[-1].hyperlink = row[-1].value
        row[-1].style = "Hyperlink"
for h in hold:
    if h.get("error"):
        excluded.append((h["symbol"], h.get("name"), "هلدینگ: " + h["error"], h.get("report_url")))

wd = wb.create_sheet("جزئیات پورتفو هلدینگ‌ها", 2)
dcols = ["هلدینگ", "شرکت سرمایه‌پذیر", "نماد", "درصد مالکیت", "تعداد سهام", "قیمت روز (ریال)",
         "ارزش روز (میلیارد تومان)", "ارزش در گزارش (میلیارد تومان)", "سال مالی گذشته", "سود خالص سال گذشته (میلیارد تومان)",
         "سال مالی جاری", "سود آینده‌نگر (میلیارد تومان)", "نسبت تقسیم سود", "سهم از سود هلدینگ (میلیارد تومان)",
         "سود نقدی دریافتی (میلیارد تومان)", "یادداشت"]
header(wd, dcols, {2: 34, 16: 50})
for h in ok:
    for d in sorted(h["detail"], key=lambda d: -(d.get("value_today") or 0)):
        wd.append([h["symbol"], d["name"], d.get("symbol"), r2(d["own"] * 100, 2), d["shares_held"], d.get("price_today"),
                   r2((d.get("value_today") or 0) / B, 1), r2((d.get("report_mv") or 0) / B, 1), d.get("inv_fy_last"),
                   r2((d.get("inv_net_last") or 0) / B, 1) if d.get("inv_net_last") is not None else None,
                   d.get("inv_fy_cur"),
                   r2((d.get("inv_net_fwd") or 0) / B, 1) if d.get("inv_net_fwd") is not None else None,
                   r2(d.get("inv_payout"), 3), r2((d.get("income_fwd") or 0) / B, 2) if d.get("income_fwd") is not None else None,
                   r2((d.get("cash_div_fwd") or 0) / B, 2) if d.get("cash_div_fwd") is not None else None, d.get("note")])

ws2 = wb.create_sheet("حذف‌شده")
header(ws2, ["نماد", "نام", "دلیل", "لینک کدال"], {2: 28, 3: 55, 4: 40})
for e in excluded:
    ws2.append(list(e))

ws3 = wb.create_sheet("داده خام (میلیون ریال)")
rcols = ["symbol", "name", "sector", "report", "months", "price", "shares", "rev", "rev_prev_same", "cogs",
         "net", "fis", "ttm_net", "q_cur", "q_prev", "receivables", "total_assets", "avg_inventory",
         "div_ratios", "period_days", "format"]
header(ws3, rcols)
for s, r in raw.items():
    ws3.append([json.dumps(r.get(k), ensure_ascii=False) if isinstance(r.get(k), list) else r.get(k) for k in rcols])

ws4 = wb.create_sheet("روش محاسبه")
ws4.sheet_view.rightToLeft = True
ws4.column_dimensions["A"].width = 130
notes = [
    "مبنا: خسروانی و آجرلو، سیستم امتیازدهی بنیادی FIS، فصلنامه چشم‌انداز حسابداری و مدیریت، دوره ۹، شماره ۱۱۶، ۱۴۰۵.",
    "داده‌ها: آخرین صورت مالی ن-۱۰ (غیرتلفیقی) هر نماد از کدال + قیمت پایانی و تعداد سهام از TSETMC در تاریخ اجرا.",
    "سود FIS = سود عملیاتی + هزینه استثنایی کاهش ارزش دریافتنی‌ها − هزینه مالی − مالیات سال جاری (حذف اقلام غیرعملیاتی و مالیات سال‌های قبل).",
    "فروش آینده‌نگر (طبق مقاله) = فروش تجمعی آخرین گزارش فعالیت ماهانه (ن-۳۰) ÷ تعداد ماه‌ها × ۱۲.",
    "سود آینده‌نگر (محافظه‌کارانه، فقط سود تحقق‌یافته) = فروش آینده‌نگر × حاشیه سود خالص آخرین صورت میان‌دوره‌ای بدون «سایر درآمدها» و سود غیرعملیاتی + سود تکرارشونده (برای هلدینگ‌ها: حاشیه سال مالی قبل). همه هزینه‌ها (سایر هزینه‌ها، زیان غیرعملیاتی) کم می‌شوند. بدون گزارش ماهانه: سود FIS میان‌دوره‌ای × ۱۲ ÷ ماه‌ها.",
    "سود غیرعملیاتی = سود قبل از مالیات (سود خالص + مالیات) − سود عملیاتی + هزینه مالی؛ سود فروش دارایی زیستی مولد (دام) جزو آن نیست و در سود می‌ماند.",
    "سود تکرارشونده: هر قلم (سایر درآمدها، سود غیرعملیاتی) که در دوره جاری و هر دو سال مالی قبل مثبت بوده، به اندازه کمترینِ این سه (دوره جاری سالانه‌شده) به برآورد سالانه اضافه می‌شود.",
    "افزایش سرمایه بعد از آخرین گزارش: محل آن از اطلاعیه کدال (ن-۵۷ تصمیمات مجمع فوق‌العاده یا ن-۷۳ تصمیمات هیئت‌مدیره) خوانده می‌شود. بخش «مطالبات و آورده نقدی» هنوز سودی در گزارش‌ها نساخته، پس در P/E، P/S، P/D و P/E TTM از ارزش بازار کم می‌شود (= ارزش بازار پیش از افزایش)؛ در هلدینگ‌ها به NAV اضافه می‌شود. افزایش از محل سود انباشته، اندوخته یا تجدید ارزیابی اثری ندارد.",
    "ROA = سود خالص ۱۲ ماه اخیر ÷ کل دارایی‌ها؛ ROE = سود خالص ۱۲ ماه اخیر ÷ حقوق صاحبان سهام (پایان دوره).",
    "دوره وصول = میانگین مطالبات (ابتدا و پایان دوره) ÷ درآمد ۱۲ ماه اخیر × ۳۶۵؛ دوره گردش کالا = میانگین موجودی ÷ بهای تمام شده ۱۲ ماه اخیر × ۳۶۵. «۱۲ ماه اخیر» یعنی دوره جاری جای همان دوره سال قبل در سال مالی قبل می‌نشیند.",
    "شرکت‌هایی که مطالبات‌شان بیش از ۸۰٪ کل دارایی‌هاست حذف می‌شوند (برگه حذف‌شده، با دلیل). هلدینگ‌ها مستثنا هستند چون مطالباتشان بیشتر سود سهام دریافتنی است.",
    "رشد فروش تجمعی و «سال مالی» (ماه‌های پوشش) از گزارش ماهانه؛ اگر ستون دوره مشابه سال قبل نداشت، از صورت میان‌دوره‌ای.",
    "تطبیق با برآورد کارشناس: تکیمیا در ۱۳ مهر ۱۴۰۵ — P/E فوروارد 2.99 و P/S فوروارد 2.05 دقیقاً بازتولید شد.",
    "حاشیه سود = سود FIS ÷ درآمد عملیاتی دوره تجمعی؛ رشد روند حاشیه = حاشیه فصل آخر نسبت به فصل قبل.",
    "نسبت تقسیم سود (فرمول مقاله) = میانگین (سود سهام مصوب ÷ سود انباشته ابتدای همان سال) برای دو مجمع اخیر؛ اگر مجمع امسال هنوز در گزارش نیست، مجمع سال قبل از آن از صورت مالی سالانه می‌آید.",
    "P/D آینده‌نگر = P/E آینده‌نگر ÷ نسبت تقسیم سود (یعنی ارزش بازار ÷ سود آینده‌نگر × نسبت تقسیم)؛ برای شرکت بدون تقسیم سود یا سود آینده‌نگر منفی خالی است. هرچه کمتر، بهتر.",
    "حذف: صندوق‌ها و شرکت‌های زیان‌ده (سود خالص ۱۲ ماه گذشته ≤ صفر) طبق مقاله امتیازدهی نمی‌شوند.",
    "نمره نرمال‌شده = امتیاز خام ÷ حداکثر امتیاز متغیرهای قابل‌محاسبه × ۱۰۰ (برای بانک/هلدینگ که بعضی متغیرها معنا ندارند). رتبه‌بندی بر اساس همین ستون است.",
    f"شرکت‌هایی که «حداکثر قابل‌محاسبه»شان کمتر از {MIN_MAX_AVAILABLE} است (بانک، بیمه و شرکت‌هایی که بخش بزرگی از متغیرها برایشان حساب نمی‌شود) در رتبه‌بندی و میانگین نیستند و در برگه «سایر شرکت‌ها» آمده‌اند.",
    f"میانگین نمره نرمال‌شده گروه: {avg:.1f} — قاعده مقاله: انتخاب از میان شرکت‌های بالاتر از میانگین.",
    "محدوده P/E خرید = ۱÷اخزا + ۱ تا ۱.۵ (با اخزا ۴۰٪: ۳.۵ تا ۴)، خروج حدود ۷ تا ۸.",
    "آستانه‌های امتیاز هر متغیر در مقاله منتشر نشده و کالیبراسیون skill برای بورس تهران است.",
    "— شرکت‌های سرمایه‌گذاری و هلدینگ‌ها (تب دوم) با روش جداگانه ارزیابی می‌شوند و در رتبه‌بندی FIS نیستند:",
    "NAV = جمع حقوق صاحبان سهام + سود تحقق‌نیافته پورتفوی بورسی (ارزش روز: تعداد سهام آخرین گزارش پورتفوی ماهانه × قیمت پایانی TSETMC، منهای بهای تمام شده انتهای دوره در همان گزارش). سرمایه‌گذاری غیربورسی فقط به بهای تمام شده (همان که در حقوق صاحبان سهام هست).",
    "شاخص اصلی = ارزش بازار ÷ (۸۰٪ × NAV)، به درصد (مثلاً ۴۰ یعنی ارزش بازار ۴۰٪ ِ ۸۰٪ NAV است)؛ هرچه کمتر، بهتر. رتبه‌بندی تب دوم بر همین اساس است.",
    "سود آینده‌نگر هلدینگ (استاندارد حسابداری): مالکیت زیر ۲۰٪ ← سود نقدی در تاریخ مجمع سرمایه‌پذیر (فقط اگر مجمع داخل سال مالی هلدینگ باشد؛ DPS واقعی اگر مجمع رفته، وگرنه EPS آینده‌نگر × نسبت تقسیم سود قبلی). مالکیت ۲۰٪ و بیشتر ← سهم از سود خالص سالی که داخل سال مالی هلدینگ تمام می‌شود.",
    "سود هر سرمایه‌پذیر با همان روش ماهانه‌ی تب FIS پیش‌بینی می‌شود (برای سرمایه‌پذیرِ هلدینگ: سود سال مالی قبل). اگر صورت مالی پیدا نشد: جدول سود سهام گزارش پورتفو یا EPS تخمینی TSETMC × میانه نسبت تقسیم.",
    "هزینه‌های هلدینگ = درآمد − سود خالص سال مالی قبل. سود تقسیمی آینده‌نگر = (سود نقدی دریافتی − هزینه‌ها) × نسبت تقسیم سود قبلی هلدینگ؛ سود تقسیمی آخرین مجمع هم جدا آمده است.",
    "این فایل خروجی یک مدل پژوهشی است و توصیه خرید یا فروش نیست.",
]
for t in notes:
    ws4.append([t])
    ws4.cell(ws4.max_row, 1).alignment = Alignment(wrap_text=True, horizontal="right")

wb.save(OUT)
print(f"saved {OUT}: ranked={len(rows)} excluded={len(excluded)} avg={avg:.1f} others={len(others)}")
for i, x in enumerate(rows[:15], 1):
    print(i, x["r"]["symbol"], x["o"]["normalized_100"], x["o"]["total"], x["o"]["max_available"], x["o"].get("pe_signal"))
